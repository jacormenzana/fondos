# proyecto3/tests/test_monthly_report_rotacion.py
# -*- coding: utf-8 -*-
"""FND-0155: the 4_Rotacion sheet renders the rotation plan persisted in the scenario notes.

Pure (R-7): a stub connection stands in for the DB.
"""
from __future__ import annotations

import json

from openpyxl import Workbook

from proyecto3.src.monthly_report import _build_rotacion, _rotation_plan_from_notes

PLAN = [
    {"subportfolio": "Defensiva", "isin_out": "ES0001", "isin_in": "ES0002",
     "recomendar": True, "coste_est": 0.0075, "razon": "Mejora 0.12 > coste"},
    {"subportfolio": "Dinamica", "isin_out": "ES0003", "isin_in": None,
     "recomendar": False, "coste_est": 0.005, "razon": "Mejora insuficiente"},
]


class _Result:
    def __init__(self, rows):
        self._rows = rows

    def fetchone(self):
        return self._rows[0] if self._rows else None

    def fetchall(self):
        return self._rows


class _Conn:
    def __init__(self, notes):
        self.notes = notes

    def execute(self, sql, params=()):
        if "portfolio_scenarios" in sql:
            return _Result([] if self.notes is None else [(self.notes,)])
        return _Result([("ES0001", "Fondo Uno"), ("ES0002", "Fondo Dos")])


def _cells(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_plan_from_notes():
    assert _rotation_plan_from_notes(json.dumps({"rotation_plan": PLAN})) == PLAN
    assert _rotation_plan_from_notes(json.dumps({"regime": "x"})) is None
    assert _rotation_plan_from_notes(None) is None
    assert _rotation_plan_from_notes("not json") is None
    assert _rotation_plan_from_notes(json.dumps([1, 2])) is None


def test_sheet_renders_plan_rows():
    ws = Workbook().active
    _build_rotacion(ws, _Conn(json.dumps({"rotation_plan": PLAN})), "esc")
    text = " | ".join(str(v) for row in _cells(ws) for v in row if v)
    assert "ES0001 Fondo Uno" in text and "ES0002 Fondo Dos" in text
    assert "0.75%" in text and "Si" in text
    assert "No (se mantiene el titular)" in text


def test_sheet_explains_when_there_is_no_plan():
    for notes in (None, json.dumps({"regime": "x"})):
        ws = Workbook().active
        _build_rotacion(ws, _Conn(notes), "esc")
        assert any("Sin plan de rotacion" in str(v) for row in _cells(ws) for v in row if v)
